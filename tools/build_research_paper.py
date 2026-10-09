from pathlib import Path
from copy import deepcopy
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf import PdfReader
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path('/Users/haru/Downloads/Example_paper_template-IJMECS (1).docx')
REVIEW = Path('/Users/haru/Documents/Codex/2026-08-02/i-have-uploaded-two-files-review/outputs/ieee_misinformation_review_with_cover_compiled.pdf')
DIAGRAM_SOURCE = Path('/Users/haru/Downloads/WhatsApp Image 2026-08-10 at 21.42.35.jpeg')
ASSETS = ROOT / 'output' / 'research-paper-assets'
OUTPUT = ROOT / 'output' / 'ai_misinformation_cognitive_bias_research_paper_draft.docx'


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)


def clear_body(document):
    body = document._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_margins(section):
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)


def add_page_field(paragraph):
    run = paragraph.add_run()
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    run._r.addnext(field)


def configure_own_header_footer(document):
    """Replace retained example-paper page furniture with project-specific text."""
    section = document.sections[0]
    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.clear()
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header_paragraph.add_run('An AI-Based Misinformation and Cognitive Bias Detection System for News Articles')
    run.font.name = 'Times New Roman'
    run.font.size = Pt(8.5)
    run.italic = True

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.clear()
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left = footer_paragraph.add_run('Taware, Mohite, Sarvade, and Sus  |  Page ')
    left.font.name = 'Times New Roman'
    left.font.size = Pt(8.5)
    add_page_field(footer_paragraph)


def add_text(document, text, style='Normal', bold=False, italic=False, align=None):
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if align is not None:
        paragraph.alignment = align
    return paragraph


def add_heading(document, number, title):
    paragraph = document.add_paragraph(style='Heading 1')
    paragraph.add_run(f'{number}. {title}')
    return paragraph


def add_subheading(document, number, title):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(f'{number}. {title}')
    run.bold = True
    return paragraph


def add_caption(document, caption):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(caption)
    run.bold = True
    run.italic = True
    run.font.size = Pt(9)
    return paragraph


def extract_review_assets():
    ASSETS.mkdir(parents=True, exist_ok=True)
    reader = PdfReader(REVIEW)
    figure_paths = []
    for page_no, page in enumerate(reader.pages, 1):
        for image_no, image in enumerate(page.images):
            suffix = Path(image.name).suffix or '.png'
            path = ASSETS / f'review_page_{page_no}_figure_{image_no}{suffix}'
            path.write_bytes(image.data)
            figure_paths.append(path)

    reference_text = '\n'.join(page.extract_text() or '' for page in reader.pages[8:])
    match = re.search(r'REFERENCES\s*(\[1\].*)', reference_text, re.S)
    entries = []
    if match:
        raw = match.group(1).replace('\u00ad', '')
        pieces = re.split(r'(?=\[\d+\])', raw)
        for piece in pieces:
            normalized = ' '.join(piece.split())
            if re.match(r'^\[\d+\]', normalized):
                entries.append(normalized)
    return figure_paths, entries


def create_corrected_architecture():
    """Create a code-native diagram that documents actual implemented modules only."""
    target = ASSETS / 'implemented_system_architecture.png'
    width, height = 1600, 2000
    image = Image.new('RGB', (width, height), 'white')
    draw = ImageDraw.Draw(image)
    try:
        title_font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial Bold.ttf', 31)
        box_font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial Bold.ttf', 23)
        body_font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 19)
        note_font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 17)
    except OSError:
        title_font = box_font = body_font = note_font = ImageFont.load_default()

    def centered(lines, y, font, fill=(20, 45, 75)):
        for line in lines:
            left, top, right, bottom = draw.textbbox((0, 0), line, font=font)
            draw.text(((width - (right - left)) / 2, y), line, font=font, fill=fill)
            y += (bottom - top) + 8

    def box(top, heading, lines, outline=(27, 73, 120), fill=(247, 250, 253), height=165):
        x1, x2 = 170, width - 170
        draw.rounded_rectangle((x1, top, x2, top + height), radius=12, outline=outline, width=3, fill=fill)
        centered([heading], top + 18, box_font)
        centered(lines, top + 58, body_font, fill=(45, 55, 70))
        return top + height

    def arrow(start, end):
        x = width / 2
        draw.line((x, start, x, end - 17), fill=(38, 58, 82), width=4)
        draw.polygon([(x - 10, end - 17), (x + 10, end - 17), (x, end)], fill=(38, 58, 82))

    centered(['SYSTEM ARCHITECTURE OF THE IMPLEMENTED PROTOTYPE'], 32, title_font)
    y = 100
    y = box(y, '1. INPUT LAYER', ['Pasted text | PDF / DOCX / TXT | Public article URL'])
    arrow(y, y + 52); y += 52
    y = box(y, '2. TEXT INGESTION, PREPROCESSING AND VALIDATION', ['URL / file text extraction | Normalisation | Size limits | Safe public URL checks'])
    arrow(y, y + 52); y += 52
    y = box(y, '3. CLAIM IDENTIFICATION', ['Sentence-level statements | Heuristic tags: factual claim, opinion, prediction, question'])
    arrow(y, y + 52); y += 52
    y = box(y, '4. LOCAL EVIDENCE RETRIEVAL', ['Local article corpus | Source and optional URL metadata', 'Keyword retrieval by default | Optional FAISS semantic retrieval'], outline=(28, 92, 130), fill=(242, 248, 253), height=180)
    arrow(y, y + 48); y += 48

    columns = [(80, 540), (570, 1030), (1060, 1520)]
    heads = [
        ('CLAIM-EVIDENCE COMPARISON', ['Relevance gate', 'Aligned / contradicted / mixed', 'Insufficient local evidence']),
        ('LANGUAGE AND FRAMING ANALYSIS', ['Emotion, certainty, authority, causality', 'Loaded labels, distrust, threat framing', 'Observable wording signals only']),
        ('SOURCE-COVERAGE VIEW', ['Group retrieved passages by source', 'Show differing retrieved emphasis', 'No ideology or quality label']),
    ]
    row_top, row_bottom = y, y + 260
    for (x1, x2), (heading, lines) in zip(columns, heads):
        draw.rounded_rectangle((x1, row_top, x2, row_bottom), radius=12, outline=(27, 73, 120), width=3, fill=(247, 250, 253))
        text_y = row_top + 22
        for line in [heading]:
            bbox = draw.textbbox((0, 0), line, font=box_font)
            draw.text((x1 + (x2-x1-(bbox[2]-bbox[0]))/2, text_y), line, font=box_font, fill=(20,45,75))
        text_y += 56
        for line in lines:
            bbox = draw.textbbox((0,0), line, font=note_font)
            draw.text((x1 + (x2-x1-(bbox[2]-bbox[0]))/2, text_y), line, font=note_font, fill=(45,55,70))
            text_y += 42
    y = row_bottom
    arrow(y, y + 50); y += 50
    y = box(y, '6. USER REPORT', ['Claims and status | Traceable excerpts and sources | Language and framing findings', 'Cross-source coverage | Clear local-dataset limitations'], height=175)
    arrow(y, y + 50); y += 50
    y = box(y, '7. USER-REVIEWED CORPUS ADDITION', ['User may save a reviewed article with source name and optional URL', 'The added article becomes searchable in the local corpus'], outline=(28, 128, 110), fill=(241, 252, 249), height=175)
    draw.rounded_rectangle((90, y + 55, width - 90, y + 165), radius=8, outline=(130, 145, 160), width=2, fill=(245, 247, 250))
    centered(['The system supports critical reading and evidence comparison. It does not determine', 'absolute truth, author intent, political ideology, or morality.'], y + 75, body_font, fill=(25, 45, 70))
    image.save(target)
    return target


def main():
    figures, references = extract_review_assets()
    architecture = create_corrected_architecture()
    document = Document(TEMPLATE)
    clear_body(document)
    set_margins(document.sections[0])
    configure_own_header_footer(document)
    styles = document.styles
    styles['Normal'].font.name = 'Times New Roman'
    styles['Normal'].font.size = Pt(10.5)
    styles['Heading 1'].font.name = 'Times New Roman'
    styles['Heading 1'].font.size = Pt(13)
    styles['Heading 1'].font.bold = True

    title = add_text(document, 'An AI-Based Misinformation and Cognitive Bias Detection System for News Articles', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    title.runs[0].font.size = Pt(17)
    title.runs[0].font.name = 'Times New Roman'

    authors = add_text(document, 'Ayush Rajshekhar Taware, Vardan Dipak Mohite, Arpita Amol Sarvade, Pranjali Jayesh Sus', align=WD_ALIGN_PARAGRAPH.CENTER)
    authors.runs[0].font.size = Pt(11)
    affiliation = add_text(document, 'Department of Computer Science and Engineering (Artificial Intelligence and Machine Learning), Rajarambapu Institute of Technology, Rajaramnagar-Ishwarpur, Maharashtra, India', align=WD_ALIGN_PARAGRAPH.CENTER)
    affiliation.runs[0].font.size = Pt(9.5)
    contact = add_text(document, 'Email: [to be added by authors]    ORCID iD: [to be added by authors]', italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    contact.runs[0].font.size = Pt(9)
    guide = add_text(document, 'Guided by Prof. Apurva A. Mohite, Assistant Professor', italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    guide.runs[0].font.size = Pt(9)

    abstract = add_text(document, '')
    abstract.add_run('Abstract: ').bold = True
    abstract.add_run(
        'This paper presents an AI-based prototype for assisting the critical reading of news articles. Rather than issuing an absolute true-or-false verdict, the system accepts article text, public URLs, and supported document files; identifies sentence-level claims; retrieves related material from a local article corpus; and reports whether the available corpus presents aligned, contradictory, mixed, or insufficient evidence. The report preserves source provenance where available and highlights observable persuasive-language and political-framing signals. The implementation combines FastAPI services, a Streamlit interface, safe text extraction, lexical retrieval with an optional semantic FAISS path, a relevance gate, explainable comparison heuristics, and source-aware reporting. The design is intentionally conservative: an absence of local evidence is reported as a coverage limitation rather than as falsity, and language signals are not treated as proof of author intent. The prototype offers a transparent workflow for evidence navigation and cognitive-bias awareness in news reading.'
    )
    index_terms = add_text(document, '')
    index_terms.add_run('Index Terms: ').bold = True
    index_terms.add_run('Misinformation Detection, Cognitive Bias, Evidence Retrieval, Explainable AI, Claim-Evidence Comparison, News Analysis.')

    add_heading(document, '1', 'Introduction')
    add_text(document, 'Digital news circulation allows reports, commentary, and user-generated material to reach large audiences with very little delay. This environment makes it difficult for readers to distinguish a factual statement, a prediction, and a strongly framed interpretation. Existing research has examined misinformation detection through machine learning, deep learning, transformer models, automated fact checking, retrieval-augmented generation, and explainable artificial intelligence [1]-[8]. However, a reader-facing system must expose the evidence and the limitations of its analysis instead of replacing the reader\'s judgment.')
    add_text(document, 'The implemented system addresses this requirement as an evidence-aware news-analysis prototype. It accepts article content through pasted text, files, or public web links. It then compares identifiable statements against a local article corpus and presents the excerpts and source information that justify each result. A separate wording analysis identifies configured cognitive and political-framing signals. The system does not classify an outlet, author, or text as morally or ideologically biased; it reports observable language patterns and corpus-limited evidence only.')
    add_text(document, 'The contribution of this paper is a transparent pipeline that combines input validation, claim identification, local retrieval, relevance filtering, claim-evidence comparison, source-aware reporting, and framing-signal analysis in one accessible interface. The remainder of the paper reviews the relevant foundations, documents the methodology and implementation, and discusses the behaviour and limits of the prototype.')

    add_heading(document, '2', 'Literature Review')
    add_text(document, 'Earlier misinformation-detection work commonly relied on supervised classifiers that used linguistic, stylistic, sentiment, readability, source, or propagation features. Classical methods remain valuable baselines, but their performance can depend heavily on labelled data and on the similarity between training and deployment domains [13], [18], [30], [58]. Neural models and transformer architectures improve contextual representation, while large language models can support claim interpretation and explanation; nevertheless, their outputs require careful grounding and transparent safeguards [21], [25], [26], [36]-[41].')
    add_text(document, 'Evidence retrieval is central to grounded analysis. RAG-oriented approaches retrieve material before producing an answer or assessment, allowing readers to inspect the passages that informed the result [19], [23], [35], [44], [60]. Benchmark resources such as FEVER, LIAR, FakeNewsNet, and VeriTeC have helped establish the broader field of claim verification [18], [29], [55]-[58]. These datasets are discussed in the literature; they are not claimed as data used by the present local-corpus prototype.')
    add_text(document, 'Research also identifies explainability and cognitive-bias awareness as important requirements. Explanations should make the basis of an assessment inspectable, while cognitive-bias and media-framing analysis should avoid inferring an author\'s intention or political identity from isolated wording [10], [12], [14], [15], [20], [32]-[34]. In response, the present system reports sentence-level language signals, source-linked evidence, and explicit insufficient-evidence outcomes.')

    if figures:
        document.add_picture(str(figures[0]), width=Inches(5.9))
        add_caption(document, 'Fig. 1. Evolution of AI-based misinformation detection (retained from the reviewed literature foundation).')

    add_heading(document, '3', 'Methodology')
    add_subheading(document, '3.1', 'Design Approach')
    add_text(document, 'The work follows a design-and-implementation approach. The aim is to construct a usable, explainable prototype for analysing news articles before sharing them. The methodology is defined by the behaviour of the implemented system: it treats the local corpus as a limited evidence resource, separates evidence comparison from language analysis, and preserves a distinction between relevance, support-like alignment, contradiction-like wording, and insufficient evidence.')
    add_subheading(document, '3.2', 'Input Acquisition and Text Preparation')
    add_text(document, 'The system accepts pasted news text, PDF, DOCX, and UTF-8 TXT files, and public HTTP/HTTPS article URLs. File extraction obtains readable text from supported formats. URL ingestion validates the scheme and host, blocks private or local addresses, limits redirects, checks content type and page size, and removes common page boilerplate before extracting article paragraphs. Submitted text is normalised by collapsing whitespace and is checked against character and word limits. These steps reduce malformed, oversized, or unsafe inputs before analysis.')
    add_subheading(document, '3.3', 'Claim Identification')
    add_text(document, 'The input is first converted into a central claim context. When an optional summarisation model is intentionally enabled and locally available, it can create a concise summary. Otherwise, the system preserves a normalised text fallback to avoid model downloads during a request. For claim-by-claim analysis, the text is segmented into up to five readable sentences. Each sentence is heuristically labelled as a factual claim, opinion, prediction, or question. This is sentence-level statement identification, not a trained atomic-claim extraction model.')
    add_subheading(document, '3.4', 'Local Evidence Retrieval and Provenance')
    add_text(document, 'The evidence resource is a local directory of article text files. Documents are divided into fixed-size word chunks. User-reviewed additions are stored with a source name and, when supplied, the original URL. Keyword retrieval is the default retrieval method and is designed to work in an offline setting. If the embedding model is explicitly enabled, sentence embeddings and a FAISS index provide an optional semantic retrieval path. Retrieved records retain the excerpt, source name, optional URL, and document identifier so that evidence shown to a reader remains traceable.')
    add_subheading(document, '3.5', 'Relevance-Gated Claim-Evidence Comparison')
    add_text(document, 'For every identified statement, the system retrieves candidate passages and applies a relevance gate before presenting any passage as evidence. Generic subject words, such as policy, government, report, and security, are excluded from the comparison vocabulary. A passage must share at least three distinctive claim terms and satisfy a claim-coverage threshold, or share at least five distinctive terms, to be considered relevant. This prevents broad topical overlap from being represented as support.')
    add_text(document, 'Relevant passages are compared using a conservative text heuristic. Compatible wording is reported as “aligned with available evidence”; directly negated shared content is reported as “contradicted by available evidence”; differing retrieved relationships are reported as “mixed evidence”; and the absence of relevant passages is reported as “insufficient evidence in the local dataset.” These labels describe the current local corpus and do not establish factual truth, credibility, consensus, or source quality.')
    add_subheading(document, '3.6', 'Cognitive-Bias and Political-Framing Signal Analysis')
    add_text(document, 'The language-analysis module applies explainable pattern matching to complete input sentences. It identifies configured signals of emotional language, absolute certainty, appeals to authority, and strong causal framing. A related module identifies loaded political labels, institutional-distrust framing, and political-threat framing. Each result contains the matched sentence, the triggering expression, and an explanation of why the wording merits verification. The modules describe textual signals only and do not diagnose the writer\'s cognitive state, ideology, intent, or morality.')
    add_subheading(document, '3.7', 'Report Generation and User-Reviewed Corpus Addition')
    add_text(document, 'The user report brings together the central claim context, the sentence-level comparison outcomes, traceable source excerpts, wording signals, source-coverage groupings, and limitations. Related passages can be grouped by their stored source label to make differences in retrieved emphasis visible. The interface also allows a reader to add a reviewed URL article to the local corpus. This is an explicit action; reviewed documents are not automatically retained after every URL analysis.')
    document.add_picture(str(architecture), width=Inches(6.2))
    add_caption(document, 'Fig. 2. System architecture of the implemented AI-based misinformation and cognitive-bias analysis prototype.')

    add_heading(document, '4', 'System Implementation and Analysis')
    add_text(document, 'The prototype is implemented as a Python application. FastAPI exposes analysis, URL-extraction, file-analysis, and knowledge-base endpoints. Streamlit provides the user interface for text, file, and URL workflows. The backend modules isolate document ingestion, claim-context extraction, local retrieval, evidence comparison, language-signal detection, political-framing detection, metrics, and explanation generation. This separation makes each analysis stage inspectable and testable.')
    add_text(document, 'The interface presents the output in reader-oriented sections. It displays the central claim context; each sentence-level status; relevant excerpts with saved source and URL information; configured cognitive and political-framing signals; a cross-source coverage view when more than one source is available; and clear limitations. If no candidate passage satisfies the relevance gate, the interface explicitly reports that no sufficiently relevant evidence was found in the local corpus.')
    add_text(document, 'Implementation checks cover document ingestion, short-input claim extraction, retrieval fallback, knowledge-base addition, source provenance, language-signal extraction, political-framing extraction, claim comparison, direct negation handling, and rejection of generic topical overlap. The project test suite contains fourteen passing tests. The system is therefore presented as a functional prototype with deterministic offline behaviour, rather than as a benchmark-trained misinformation classifier.')

    add_heading(document, '5', 'Discussion')
    add_text(document, 'The prototype is designed to make its reasoning inspectable. Its principal output is not a binary misinformation label; it is a structured account of what the local evidence corpus contains, what wording patterns occur in the submitted article, and what remains uncertain. This framing reduces the risk that weak keyword matches, unsupported model claims, or the absence of local evidence are interpreted as definitive truth judgments.')
    add_text(document, 'The relevance gate is particularly important for the local-corpus setting. Without it, generic vocabulary can retrieve broadly related but non-evidential passages. By requiring distinctive shared content before a passage is shown as evidence, the system gives priority to “insufficient evidence” over an unsupported alignment claim. Its limitation is equally important: outputs depend on the coverage and provenance of the articles stored in the local corpus. A result is therefore an evidence-navigation aid, not an independent verification authority.')

    add_heading(document, '6', 'Conclusion')
    add_text(document, 'This paper documented an AI-based system for misinformation and cognitive-bias analysis in news articles. The implementation accepts multiple input forms, identifies sentence-level statements, retrieves and filters source-traceable local evidence, reports cautious comparison outcomes, and exposes configured cognitive and political-framing signals. By clearly separating corpus evidence, wording signals, and limitations, the prototype supports critical reading without claiming to determine absolute truth, author intent, ideology, or morality.')

    add_heading(document, '7', 'All the Declarations and Statements')
    add_subheading(document, '7.1', 'Author Contributions Statement')
    add_text(document, 'Ayush Rajshekhar Taware, Vardan Dipak Mohite, Arpita Amol Sarvade, and Pranjali Jayesh Sus contributed to the project development, literature synthesis, implementation, testing, and manuscript preparation. Prof. Apurva A. Mohite provided academic guidance and supervision. [To be finalised by the authors.]')
    add_subheading(document, '7.2', 'Conflict of Interest Statement')
    add_text(document, 'The authors declare no conflict of interest. [To be confirmed by the authors.]')
    add_subheading(document, '7.3', 'Data Availability Statement')
    add_text(document, 'The prototype operates on the local article corpus bundled with the project and on user-reviewed articles explicitly added through the interface. [To be finalised according to conference requirements.]')

    add_heading(document, '8', 'References')
    for entry in references:
        add_text(document, entry)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)
    print(f'REFERENCES={len(references)} FIGURES={len(figures)}')


if __name__ == '__main__':
    main()
